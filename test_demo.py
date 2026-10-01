# -*- coding: utf-8 -*-
# Demo 功能专项测试（不触网）
import importlib.util, os, sys, tempfile

demo_path = r"C:\Users\TianYifan\.doubao\lark-chats\2026-10-01\new-chat\re_work\jira_detail_export_demo.py"
spec = importlib.util.spec_from_file_location("demo_mod", demo_path)
mod = importlib.util.module_from_spec(spec)
sys.modules["demo_mod"] = mod
spec.loader.exec_module(mod)

DemoJiraClient = mod.DemoJiraClient
write_excel = mod.write_excel

# ---- 样例：改动记录 tab HTML（仿 HAR 结构）----
CHANGE_HISTORY_HTML = """
<div class="issue-data-block" id="changehistory-20527936">
    <div class="actionContainer">
        <div class="action-details" id="changehistorydetails_20527936">
    <a class="user-hover user-avatar" rel="vlxauh" id="changehistoryauthor_20527936" href="/secure/ViewProfile.jspa?name=vlxauh"><span class="aui-avatar aui-avatar-xsmall"><span class="aui-avatar-inner"><img src="http://jira.z-onesoftware.com:8080/secure/useravatar?size=xsmall&amp;avatarId=10122" alt="vlxauh" /></span></span> Li Ziwen 李子文</a>
进行了改变  - <span class='date' title='2026/08/02 08:58'><time class='livestamp' datetime='2026-08-02T08:58:15+0800'>2026/08/02 08:58</time></span>
            </div>
        <div class="changehistory action-body">
<table class="wiki-content">
<tr><th>域</th><th>原值</th><th>新值</th></tr>
<tr><td>经办人</td><td></td><td>Qin Yanan 秦亚南 [JIRAUSER34739]</td></tr>
</table>
        </div>
    </div>
</div>
<div class="issue-data-block" id="changehistory-20527947">
    <div class="actionContainer">
        <div class="action-details" id="changehistorydetails_20527947">
    <a class="user-hover user-avatar" rel="gucviw" id="changehistoryauthor_20527947" href="/secure/ViewProfile.jspa?name=gucviw"><span class="aui-avatar aui-avatar-xsmall"><span class="aui-avatar-inner"><img src="http://jira.z-onesoftware.com:8080/secure/useravatar?size=xsmall&amp;avatarId=10334" alt="gucviw" /></span></span> Luo Changchuan 罗常钏</a>
进行了改变  - <span class='date' title='2026/08/05 09:47'><time class='livestamp' datetime='2026-08-05T09:47:17+0800'>2026/08/05 09:47</time></span>
            </div>
        <div class="changehistory action-body">
<table class="wiki-content">
<tr><th>域</th><th>原值</th><th>新值</th></tr>
<tr><td>CR编号</td><td>3ER(5s Export)</td><td>3EV(5SExport)</td></tr>
</table>
        </div>
    </div>
</div>
"""

client = DemoJiraClient("http://jira.test:8080", "u", "p")
acts = client.parse_changehistory(CHANGE_HISTORY_HTML)

assert len(acts) == 2, len(acts)
assert acts[0]["author"] == "Li Ziwen 李子文", acts[0]["author"]
assert acts[0]["time"] == "2026-08-02T08:58:15+0800", acts[0]["time"]
assert "经办人 |  | Qin Yanan 秦亚南 [JIRAUSER34739]" in acts[0]["changes"], acts[0]["changes"]
assert "CR编号 | 3ER(5s Export) | 3EV(5SExport)" in acts[1]["changes"], acts[1]["changes"]
print("PASS parse_changehistory (2 blocks, author/time/changes)")

# 表头行（域/原值/新值）被跳过
assert "域 | 原值 | 新值" not in acts[0]["changes"]
print("PASS header row skipped")

# 空块过滤（只有 issuecreated 时）
only_created = '<div class="issue-data-block" id="issuecreated-1655626"><div class="actionContainer"><div class="action-details">creator</div></div></div>'
assert client.parse_changehistory(only_created) == []
print("PASS issuecreated block ignored")

# ---- 第一个 KEY 先全部后改动记录：请求顺序模拟 ----
class FakeSession:
    def __init__(self):
        self.urls = []
    def open(self, req, timeout=None):
        self.urls.append(req.full_url)
        class R:
            def read(self):
                return b"<html/>"
        return R()

fc = DemoJiraClient("http://jira.test:8080", "u", "p")
fc.session = FakeSession()
# 手动复刻 _run_once 中首个 KEY 的逻辑
key = "BSUVVW-30904"
all_html = fc.fetch_activity_tab(key, DemoJiraClient.TAB_ALL)
ch_html = fc.fetch_activity_tab(key, DemoJiraClient.TAB_CHANGEHISTORY)
urls = fc.session.urls
assert len(urls) == 2, urls
assert "page=com.atlassian.jira.plugin.system.issuetabpanels:all-tabpanel" in urls[0], urls[0]
assert "page=com.atlassian.jira.plugin.system.issuetabpanels:changehistory-tabpanel" in urls[1], urls[1]
print("PASS first-key order: all-tabpanel then changehistory-tabpanel")

# ---- Excel 写入 ----
tmpdir = tempfile.mkdtemp()
xlsx = os.path.join(tmpdir, "test_out.xlsx")
headers = ["Issue Key", "概要", "改动记录"]
rows = [["BSUVVW-30904", "测试概要",
         "[2026-08-02T08:58:15+0800] Li Ziwen 李子文\n经办人 |  | Qin Yanan"]]
write_excel(xlsx, headers, rows)
assert os.path.exists(xlsx) and os.path.getsize(xlsx) > 1000, "xlsx too small"
from openpyxl import load_workbook
wb = load_workbook(xlsx)
ws = wb.active
assert ws["A1"].value == "Issue Key"
assert ws["B2"].value == "测试概要"
assert "Li Ziwen" in ws["C2"].value
print("PASS excel write (headers + data + multiline cell)")

print("\nALL DEMO TESTS PASSED")
